"""Read an explicitly mapped XLSX roster; never create/edit a workbook."""
import json, sys
from openpyxl import load_workbook

config = json.loads(open(sys.argv[1], encoding='utf-8-sig').read())['roster']
wb = load_workbook(config['path'], read_only=True, data_only=True)
records = []
try:
    for block in config['blocks']:
        sheet = wb[block['sheet']]
        for row in range(block['firstRow'], block['lastRow'] + 1):
            number = sheet.cell(row, block['numberColumn']).value
            name = sheet.cell(row, block['nameColumn']).value
            if number is None and name is None:
                continue
            if not name or number is None:
                raise ValueError('Incomplete roster row')
            if block.get('idColumn'):
                student_id = str(sheet.cell(row, block['idColumn']).value).strip()
            else:
                if int(number) != number:
                    raise ValueError('Non-integer roster number')
                student_id = str(block['idBase'] + int(number))
            if not student_id.isdigit():
                raise ValueError('Invalid roster ID')
            records.append({'id': student_id, 'name': str(name).strip()})
finally:
    wb.close()
sys.stdout.reconfigure(encoding='utf-8')
print(json.dumps(records, ensure_ascii=False))
